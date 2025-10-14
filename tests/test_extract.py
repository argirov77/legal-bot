from __future__ import annotations

import subprocess
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos

from app.extract import extract_text


def test_extract_text_from_txt(tmp_path):
    sample_path = tmp_path / "sample.txt"
    sample_path.write_text("Sample text content for extraction.")

    text = extract_text(sample_path)
    assert text.strip() and "sample text" in text.lower()


def test_extract_text_from_docx(tmp_path):
    sample_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Sample DOCX file for extraction.")
    document.save(sample_path)

    text = extract_text(sample_path)
    assert text.strip() and "docx file" in text.lower()


def test_extract_text_from_pdf(tmp_path):
    sample_path = tmp_path / "sample.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "Sample PDF fixture text.", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.output(str(sample_path))

    text = extract_text(sample_path)
    assert text.strip() and "pdf fixture" in text.lower()


def test_extract_text_pdf_triggers_ocr(tmp_path):
    pdf_path = tmp_path / "empty.pdf"
    pdf_path.write_bytes(b"")

    with patch("app.extract.pdf_extract_text", return_value="") as mock_pdf_extract:
        def run_side_effect(cmd, **kwargs):
            sidecar_index = cmd.index("--sidecar") + 1
            sidecar_path = Path(cmd[sidecar_index])
            sidecar_path.write_text("OCR Result")
            return subprocess.CompletedProcess(cmd, 0)

        with patch("app.extract.subprocess.run", side_effect=run_side_effect) as mock_run:
            text = extract_text(pdf_path, use_ocr=True)

    mock_pdf_extract.assert_called_once()
    mock_run.assert_called_once()
    assert "OCR Result" in text
